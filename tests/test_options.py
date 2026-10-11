import pytest

from global_investment_research import options as O


def test_black_scholes_known_values():
    # Hull textbook values: S=K=100, T=1, r=5%, sigma=20%
    assert O.bs_price(100, 100, 1, 0.2, "C", r=0.05) == pytest.approx(10.4506, abs=1e-4)
    assert O.bs_price(100, 100, 1, 0.2, "P", r=0.05) == pytest.approx(5.5735, abs=1e-4)
    # put-call parity
    c, p = O.bs_price(120, 100, 0.5, 0.3, "C"), O.bs_price(120, 100, 0.5, 0.3, "P")
    assert c - p == pytest.approx(120 - 100 * pytest.importorskip("math").exp(-O.RISK_FREE * 0.5))


def test_quote_mid_rejects_wide_or_empty():
    assert O.quote_mid({"bid": 1.0, "ask": 1.1})[0] == pytest.approx(1.05)
    assert O.quote_mid({"bid": 0.5, "ask": 1.5})[0] is None
    assert O.quote_mid({"bid": None, "ask": None, "last": 2.0}) == (2.0, "last trade (no two-sided quote)")
    assert O.quote_mid({})[0] is None


def test_spread_analytics():
    a = O.analyse("Bull call spread", [(1, "C", 100), (-1, "C", 110)], 4.0, 100, 0.25, 0.3)
    assert a["breakeven"] == 104 and a["max_gain"] == 6 and a["reward_risk"] == pytest.approx(1.5)
    assert a["scenarios"]["+20%"] == pytest.approx(1.5)   # capped at width
    assert a["scenarios"]["-10%"] == pytest.approx(-1.0)  # loses whole debit
    assert 0 < a["prob_profit"] < 0.5
    p = O.analyse("Bear put spread", [(1, "P", 100), (-1, "P", 90)], 3.0, 100, 0.25, 0.3)
    assert p["breakeven"] == 97 and p["scenarios"]["-20%"] == pytest.approx(7 / 3)


def test_ideas_follow_desk_and_size_to_budget():
    chain = {"underlyings": [{"symbol": "AAA", "last": 100.0, "expiries": [{"expiry": "2026-12-18", "contracts": [
        {"right": "C", "strike": k, "bid": b, "ask": b * 1.05, "iv": 0.30, "open_interest": 500}
        for k, b in [(100, 6.0), (105, 3.8), (110, 2.3)]]}]},
        {"symbol": "BBB", "last": 50.0, "expiries": [{"expiry": "2026-12-18", "contracts": []}]}]}
    dec = {"calls": [{"symbol": "AAA", "stance": "Accumulate"}, {"symbol": "BBB", "stance": "Watch"}]}
    out = O.build_ideas(chain, dec, "2026-10-09", portfolio=1_000_000)
    names = {i["structure"] for i in out["ideas"]}
    assert len(names) == 1 and names <= {"Bull call spread", "Long call"}   # one idea per underlying in the sleeve
    assert any(s["symbol"] == "AAA" and "better-ranked" in s["reason"] for s in out["skipped"])
    for i in out["ideas"]:
        assert i["premium_at_risk"] <= O.PER_TRADE_RISK * 1_000_000
    assert any(s["symbol"] == "BBB" for s in out["skipped"])
    from global_investment_research.options_desk import options_page
    html = options_page(out)
    assert "<title>Speculative Options Sleeve</title>" in html and "svg" in html and "Defined risk only" in html
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    from global_investment_research.workbook import build_workbook
    import tempfile, pathlib
    p = pathlib.Path(tempfile.mkdtemp()) / "b.xlsx"
    build_workbook({"as_of": "2026-10-09", "posture": {"risk_budget": 1, "risk_note": "", "tilts": []}, "calls": [], "fx": []},
                   p, opt=out)
    ws = load_workbook(p)["Options"]
    assert ws["O2"].value == "=INT(Policy!$B$2*Policy!$B$11/H2)"


def test_digest_email_and_telegram():
    import json as _j
    from global_investment_research import notify
    from global_investment_research import decisions as D
    from global_investment_research.cli import synthetic
    from global_investment_research.data import load_universe
    from global_investment_research.monitor import build_monitor
    px, meta = synthetic(load_universe())
    mon = build_monitor(px, meta)
    dec = D.build_decisions(mon, meta)
    d = notify.digest(dec, mon, None, [{"topic": "Rates", "from": "a", "to": "b"}])
    tg, em = notify.telegram_text(d), notify.email_html(d)
    assert "Rates: a → b" in tg and len(tg) < 4096  # Telegram message limit
    assert "<table" in em and "Weekly Investment Brief" in em

    class Resp:
        def __init__(self, ok): self.ok = ok
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def read(self): return _j.dumps({"ok": self.ok, "description": "bad"}).encode()

    sent = {}
    def opener(req, timeout):
        sent["url"], sent["data"] = req.full_url, req.data
        return Resp(True)
    notify.send_telegram("hi", "TOKEN", "123", opener=opener)
    assert sent["url"].endswith("/botTOKEN/sendMessage") and b"chat_id=123" in sent["data"]
    with pytest.raises(RuntimeError):
        notify.send_telegram("hi", "TOKEN", "123", opener=lambda r, timeout: Resp(False))


def test_earnings_note_never_invents_a_date():
    from datetime import date
    from global_investment_research.options import earnings_note
    d0 = date(2026, 10, 9)
    assert earnings_note("sector_proxy", 70, d0, "2026-12-18", None, "XLE")["status"] == "none"
    assert earnings_note("stock", 70, d0, "2026-12-18", None, "XOM")["status"] == "unverified"
    assert earnings_note("stock", 100, d0, "2027-01-17", None, "XOM")["status"] == "spans"
    n = earnings_note("stock", 70, d0, "2026-12-18", {"XOM": "2026-10-30"}, "XOM")
    assert n["status"] == "spans" and n["date"] == "2026-10-30"
    assert earnings_note("stock", 70, d0, "2026-12-18", {"XOM": "2027-01-30"}, "XOM")["status"] == "clear"


def test_sleeve_cap_is_enforced_in_the_output():
    import json
    from pathlib import Path
    from global_investment_research import options as O
    root = Path(__file__).resolve().parents[1]
    chains = json.loads((root / "data/snapshots/2026-10-11/options.json").read_text())
    dec = json.loads((root / "reports/2026-10-11/decisions.json").read_text())
    res = O.build_ideas(chains, dec, "2026-10-09")
    syms = [i["symbol"] for i in res["ideas"]]
    assert len(syms) == len(set(syms))                                   # one idea per underlying
    assert sum(i["premium_at_risk"] for i in res["ideas"]) <= O.SLEEVE_CAP * 1_000_000 + 1e-6
    assert res["sleeve_used"] <= O.SLEEVE_CAP



def test_flagged_contract_screens_out_the_idea_and_payoff_matches():
    contracts = [{"right": "C", "strike": k, "bid": b, "ask": b * 1.05, "iv": 0.30, "open_interest": 500}
                 for k, b in [(100, 6.0), (105, 3.8), (110, 2.3)]]
    chain = {"retrieved_at_utc": "2026-10-09T20:00:00Z",
             "underlyings": [{"symbol": "AAA", "last": 100.0, "expiries": [{"expiry": "2026-12-18", "contracts": contracts}]}]}
    dec = {"calls": [{"symbol": "AAA", "stance": "Accumulate"}]}
    out = O.build_ideas(chain, dec, "2026-10-09")
    assert out["chain_issues"] == [] and out["ideas"]
    contracts[2]["bid"] = 0.0                                   # 110 call: ask with no bid
    out = O.build_ideas(chain, dec, "2026-10-09")
    assert any(q["strike"] == 110 for q in out["chain_issues"])
    assert all(110 not in [l["strike"] for l in i["legs"]] for i in out["ideas"])
    assert any("chain check failed" in s["reason"] and "110C" in s["reason"] for s in out["skipped"])
